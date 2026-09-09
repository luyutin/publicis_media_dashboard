"""Exact aliases plus optional semantic mapping of structurally detected regions."""

import json
import math
import re
import urllib.request
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .grid import clean_text, normalize, is_blank, is_summary, date_parts
from .models import Candidate, OllamaConfig
from .settings import ALIASES, TARGET_COLUMNS, TARGET_DESCRIPTIONS, OLLAMA_SYSTEM_PROMPT, render_ollama_user_prompt


def read_dictionary(path: Path | None) -> dict[str, set[str]]:
    aliases = {target: {normalize(a) for a in values} for target, values in ALIASES.items()}
    if not path or not path.exists():
        return aliases
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["字典欄位說明"] if "字典欄位說明" in wb.sheetnames else wb.worksheets[0]
        canonical = {normalize(target): target for target in TARGET_COLUMNS}
        ignored = {"na", "∅", "檔名", "品牌自填", "檔名[0]", "檔名[2]", "檔名[3]"}
        for row in ws.iter_rows(min_row=3, values_only=True):
            if len(row) < 3:
                continue
            target = canonical.get(normalize(row[2]))
            if target:
                for value in (row[1], row[2], *row[7:12]):
                    alias = normalize(value)
                    if alias and alias not in ignored and not any(
                        alias in names for key, names in aliases.items() if key != target
                    ):
                        aliases[target].add(alias)
    finally:
        wb.close()
    return aliases


def excluded_header(header: str) -> bool:
    return any(is_summary(part) for part in header.split(" | "))


def match_header(header, aliases) -> str | None:
    """No substring matching: 'Cost per result' must not become total spend."""
    if excluded_header(clean_text(header)):
        return None
    parts = [clean_text(header), *reversed(clean_text(header).split(" | "))]
    for part in dict.fromkeys(parts):
        key = normalize(part)
        canonical = [t for t in TARGET_COLUMNS if normalize(t) == key]
        if canonical:
            return canonical[0]
        targets = [target for target, names in aliases.items() if key and key in names]
        if len(targets) == 1:
            return targets[0]
    return None


def plausible_mapping(header: str, target: str) -> bool:
    """Reject measurable contradictions, without requiring alias membership."""
    text = header.split(" | ")[-1].casefold()
    if target == "Spent (TWD)" and (re.search(r"\b(usd|hkd|jpy|eur|cpm|cpc|cpa)\b", text) or "per " in text):
        return False
    seconds = re.search(r"(\d+)\s*(?:s\b|sec|second|秒|\")", text)
    if seconds and ("Views" in target or "Video played" in target or target == "MV>5"):
        expected = {'3" Video Views': "3", '15" Video Views (ThruPlays)': "15"}
        if expected.get(target) != seconds.group(1):
            return False
    if target.startswith("Video played to "):
        percent = re.search(r"(\d+)\s*%", text)
        if percent and percent.group(1) not in target:
            return False
    if target in {"Impressions", "Clicks (all)", "Views", "Reach"}:
        if any(marker in text for marker in ("rate", "ctr", "cpm", "cpc", "率", "per ")):
            return False
    return True


def rule_mapping(grid, block: Candidate, aliases):
    suggestions = {}
    if block.shared_date_column is not None:
        suggestions[block.shared_date_column] = "Date"
    for col in range(block.start_column, block.end_column + 1):
        header = block.headers[col - 1]
        if excluded_header(header):
            block.excluded_columns.append(f"{col}: {header}")
            continue
        target = match_header(header, aliases)
        if target and plausible_mapping(header, target):
            suggestions[col] = target
    # Infer unfamiliar date labels from actual date values, never plain numbers.
    if "Date" not in suggestions.values():
        date_columns = []
        for col in range(block.start_column, block.end_column + 1):
            if excluded_header(block.headers[col - 1]) or not block.headers[col - 1]:
                continue
            values = [grid.get(r, col) for r in range(block.header_row + 1, min(block.end_row, block.header_row + 8) + 1)]
            filled = [v for v in values if not is_blank(v) and not is_summary(v)]
            if filled and sum(date_parts(v) is not None for v in filled) / len(filled) >= 0.6:
                date_columns.append(col)
        date_columns = [col for index, col in enumerate(date_columns) if not any(
            _shared_merged_column(grid, block, previous, col) for previous in date_columns[:index]
        )]
        if len(date_columns) == 1 and date_columns[0] not in suggestions:
            suggestions[date_columns[0]] = "Date"
        elif len(date_columns) > 1:
            block.warnings.append("多個日期欄位，需確認報表日期欄位")
    for target in dict.fromkeys(suggestions.values()):
        columns = [col for col, mapped in suggestions.items() if mapped == target]
        if target == "Spent (TWD)" and len(columns) > 1:
            explicit = [col for col in columns if re.search(r"twd|ntd|台幣|新台幣", block.headers[col - 1], re.I)]
            if len(explicit) == 1:
                block.warnings.append("花費使用明確標示 TWD/NTD 的欄位；未採用其他花費欄")
                columns = explicit
        if len(columns) == 1 or all(_shared_merged_column(grid, block, columns[0], col) for col in columns[1:]):
            block.mapped[columns[0]] = target
            for col in columns[1:]:
                block.excluded_columns.append(f"{col}: {block.headers[col - 1]}（同一合併欄位）")
        else:
            warning = f"多欄對應 {target}，未自動合併或加總"
            if warning not in block.warnings:
                block.warnings.append(warning)


def _shared_merged_column(grid, block, left, right):
    seen = False
    for row in range(block.header_row + 1, block.end_row + 1):
        if is_blank(grid.get(row, left)) and is_blank(grid.get(row, right)):
            continue
        seen = True
        if grid.origins.get((row, left), (row, left)) != grid.origins.get((row, right), (row, right)):
            return False
    return seen


def _compact(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value[:160] if isinstance(value, str) else value


def block_option(grid, block):
    date_col = next((c for c, target in block.mapped.items() if target == "Date"), block.shared_date_column)
    return {
        "id": block.block_id,
        "header_row": block.header_row,
        "start_column": block.start_column,
        "end_column": block.end_column,
        "end_row": block.end_row,
        "headers": [{"column": c, "value": block.headers[c - 1]} for c in range(block.start_column, block.end_column + 1)],
        "sample_rows": [
            {"row": r, "values": [_compact(v) for v in grid.row(r, block.start_column, block.end_column)]}
            for r in range(block.header_row + 1, min(block.end_row, block.header_row + 4) + 1)
        ],
        "context": [[_compact(v) for v in row] for row in grid.context(block.header_start, block.start_column, block.end_column)],
        "known_mappings": block.mapped,
        "shared_date_column": block.shared_date_column,
        "date_source": {
            "column": date_col,
            "header": block.headers[date_col - 1],
            "samples": [[r, _compact(grid.get(r, date_col))]
                        for r in range(block.header_row + 1, min(block.end_row, block.header_row + 4) + 1)],
        } if date_col is not None else None,
    }


def call_ollama_json(prompt: str, config: OllamaConfig, *, schema=None, system_prompt=None) -> dict:
    schema = schema or {
        "type": "object", "required": ["blocks"],
        "properties": {"blocks": {"type": "array", "items": {
            "type": "object", "required": ["id", "mappings"],
            "properties": {
                "id": {"type": "string"},
                "mappings": {"type": "array", "items": {
                    "type": "object", "required": ["column", "target", "confidence"],
                    "properties": {
                        "column": {"type": "integer", "minimum": 1},
                        "target": {"type": "string", "enum": TARGET_COLUMNS},
                        "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                    },
                }},
            },
        }}},
    }
    payload = {
        "model": config.model, "stream": False, "think": False,
        "format": schema, "options": {"temperature": 0},
        "messages": [{"role": "system", "content": system_prompt or OLLAMA_SYSTEM_PROMPT}, {"role": "user", "content": prompt}],
    }
    request = urllib.request.Request(
        f"{config.url.rstrip('/')}/api/chat", data=json.dumps(payload, ensure_ascii=False).encode(),
        headers={"Content-Type": "application/json"}, method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=config.timeout) as response:
            result = json.loads(response.read().decode())
        parsed = json.loads(result["message"]["content"])
        if not isinstance(parsed, dict) or not isinstance(parsed.get("blocks"), list):
            raise ValueError("回應必須包含 blocks 陣列")
        return parsed
    except (OSError, ValueError, TypeError, KeyError) as exc:
        raise RuntimeError(
            f"Ollama 請求失敗（model={config.model}, "
            f"url={config.url.rstrip('/')}/api/chat, {type(exc).__name__}）：{exc}"
        ) from exc


def map_blocks(grid, blocks, aliases, config: OllamaConfig | None):
    for block in blocks:
        rule_mapping(grid, block, aliases)
    # Diagnostic force_all mode sends every detected region, including regions
    # already mapped by aliases and regions later excluded for lacking dates.
    pending = list(blocks) if config and config.force_all else [b for b in blocks if _needs_llm(grid, b)]
    if config and config.enabled:
        for offset in range(0, len(pending), 6):
            batch = pending[offset:offset + 6]
            prompt = render_ollama_user_prompt(
                target_descriptions=json.dumps(TARGET_DESCRIPTIONS, ensure_ascii=False),
                options=json.dumps([block_option(grid, b) for b in batch], ensure_ascii=False),
            )
            try:
                response = call_ollama_json(prompt, config)
                if not isinstance(response.get("blocks"), list):
                    raise RuntimeError("Ollama 回應缺少 blocks 陣列，已改用規則")
                by_id = {b.block_id: b for b in batch}
                responded_ids = set()
                for item in response["blocks"]:
                    if not isinstance(item, dict) or not isinstance(item.get("id"), str):
                        continue
                    block = by_id.get(item["id"])
                    if block is None or not isinstance(item.get("mappings"), list):
                        continue
                    responded_ids.add(block.block_id)
                    block.method = "llm-layout+llm-mapping" if block.method.startswith("llm-layout") else "rules+llm"
                    proposed = {}
                    for mapping in item["mappings"]:
                        if not isinstance(mapping, dict):
                            continue
                        col, target, confidence = mapping.get("column"), mapping.get("target"), mapping.get("confidence")
                        if (type(col) is not int or not isinstance(target, str)
                                or not isinstance(confidence, (int, float)) or not math.isfinite(confidence)
                                or not 0.8 <= confidence <= 1 or target not in TARGET_COLUMNS
                                or not block.start_column <= col <= block.end_column):
                            continue
                        if any(value.startswith(f"{col}:") for value in block.excluded_columns):
                            continue
                        header = block.headers[col - 1]
                        if not header or excluded_header(header) or not plausible_mapping(header, target):
                            continue
                        if target == "Date":
                            values = [grid.get(r, col) for r in range(block.header_row + 1, min(block.end_row, block.header_row + 8) + 1)]
                            filled = [v for v in values if not is_blank(v) and not is_summary(v)]
                            if not filled or sum(date_parts(v) is not None for v in filled) / len(filled) < 0.6:
                                block.warnings.append("LLM 日期對應缺少實際日期樣本，已忽略")
                                continue
                        if col not in block.mapped and target not in block.mapped.values():
                            proposed.setdefault(col, set()).add(target)
                    for col, targets in proposed.items():
                        if len(targets) != 1:
                            block.warnings.append(f"LLM 對欄 {col} 提出衝突對應，已忽略")
                            continue
                        target = next(iter(targets))
                        if sum(target in values for values in proposed.values()) != 1:
                            block.warnings.append(f"LLM 多欄對應 {target}，已忽略")
                            continue
                        block.mapped[col] = target
                for block in batch:
                    if block.block_id not in responded_ids:
                        block.method = (
                            "llm-layout+llm-mapping-incomplete"
                            if block.method.startswith("llm-layout") else "rules+llm-incomplete"
                        )
                        block.warnings.append(
                            f"LLM 欄位辨識回應未包含區塊 {block.block_id}（sheet={grid.sheet}）"
                        )
            except RuntimeError as exc:
                config.failure = str(exc)
                block_ids = "、".join(block.block_id for block in batch)
                warning = (
                    f"LLM 欄位辨識失敗（sheet={grid.sheet}, blocks={block_ids}）："
                    f"{exc}；僅本批使用規則結果，後續批次仍會繼續呼叫 Ollama"
                )
                for block in batch:
                    block.method = (
                        "llm-layout+llm-mapping-failed"
                        if block.method.startswith("llm-layout") else "rules+llm-failed"
                    )
                    block.warnings.append(warning)
    for block in blocks:
        excluded = {int(value.split(":", 1)[0]) for value in block.excluded_columns}
        block.unmapped = [block.headers[c - 1] for c in range(block.start_column, block.end_column + 1)
                          if c not in block.mapped and c not in excluded and block.headers[c - 1]]
        width = block.end_column - block.start_column + 1
        if block.shared_date_column is not None and not block.start_column <= block.shared_date_column <= block.end_column:
            width += 1
        block.score = round(len(block.mapped) / max(1, width), 3)


def _needs_llm(grid, block):
    if not _eligible_for_llm(grid, block):
        return False
    unsupported = {"ctr", "ctr(%)", "cpm", "cpc", "cpv", "cpoc", "cpe", "vtr",
                   "點擊率", "點擊率(%)", "engagementrate", "costperresult", "costperresulttype"}
    for col in range(block.start_column, block.end_column + 1):
        header = block.headers[col - 1]
        if (not header or col in block.mapped or excluded_header(header)
                or normalize(header.split(" | ")[-1]) in unsupported
                or any(value.startswith(f"{col}:") for value in block.excluded_columns)):
            continue
        if any(not is_blank(grid.get(r, col)) for r in range(block.header_row + 1, min(block.end_row, block.header_row + 8) + 1)):
            return True
    return False


def _eligible_for_llm(grid, block):
    if is_summary(block.group_label):
        return False
    # A non-date demographic table cannot be repaired by inventing a date axis.
    return "Date" in block.mapped.values() or any(
        date_parts(grid.get(r, c))
        for r in range(block.header_row + 1, min(block.end_row, block.header_row + 8) + 1)
        for c in range(block.start_column, block.end_column + 1)
    )
