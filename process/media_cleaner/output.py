"""Excel presentation and consolidation; independent of table recognition."""

import json
from dataclasses import asdict, fields
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .grid import clean_text, is_blank
from .models import AuditRecord
from .settings import SOURCE_COLUMN, TARGET_COLUMNS

def populated_columns(records: Iterable[dict[str, Any]]) -> list[str]:
    """Template fields first, source metadata next, original fields last."""
    rows = list(records)
    declared = {column for record in rows for column in record}
    canonical = [
        column
        for column in TARGET_COLUMNS
        if column in declared
    ]
    metadata = [SOURCE_COLUMN] if SOURCE_COLUMN in declared else []
    original = []
    for record in rows:
        for column in record:
            if column not in TARGET_COLUMNS and column != SOURCE_COLUMN and column not in original:
                original.append(column)
    return canonical + metadata + original


def format_output(
    ws: Any,
    row_count: int,
    output_columns: Iterable[str] | None = None,
) -> None:
    """Apply a small, consistent format to the cleaned output worksheet."""
    columns = list(output_columns or (cell.value for cell in ws[1] if cell.value))
    if not columns:
        return
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    last_column = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_column}{max(row_count + 1, 1)}"

    widths = {
        "Date": 13,
        "Campaign name": 30,
        "Adset name": 30,
        "Ad Free Form": 30,
        "Final URL": 40,
        "Impressions": 14,
        "Clicks (all)": 14,
        "Spent (TWD)": 16,
        "Campaign Type": 18,
        SOURCE_COLUMN: 36,
    }
    column_letters = {
        target: get_column_letter(index)
        for index, target in enumerate(columns, start=1)
    }
    for target, column_letter in column_letters.items():
        ws.column_dimensions[column_letter].width = widths.get(target, 18)

    if "Date" in column_letters:
        for cell in ws[column_letters["Date"]][1:]:
            cell.number_format = "yyyy-mm-dd"
    for target in ("Bounce Rate", "TVR", "10 Second TVR"):
        if target in column_letters:
            for cell in ws[column_letters[target]][1:]:
                cell.number_format = "0.00%"
    for target in (
        "Reach",
        "Impressions",
        "Clicks (all)",
        "Link clicks (Web Clicks)",
        "Views",
        '3" Video Views',
        '15" Video Views (ThruPlays)',
        "TrueView: Views",
        "Video played to 25%",
        "Video played to 50%",
        "Video played to 75%",
        "Video played to 100%",
    ):
        if target in column_letters:
            for cell in ws[column_letters[target]][1:]:
                cell.number_format = "#,##0"


def _format_audit_sheet(
    ws: Any,
    row_count: int,
    widths: dict[str, int],
) -> None:
    """Apply compact, consistent formatting to an audit worksheet."""
    header_fill = PatternFill("solid", fgColor="5B6573")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    last_column = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{last_column}{max(row_count + 1, 1)}"

    header_columns = {
        cell.value: get_column_letter(index)
        for index, cell in enumerate(ws[1], start=1)
    }
    for header, width in widths.items():
        column = header_columns.get(header)
        if column:
            ws.column_dimensions[column].width = width


def consolidate_cleaned_workbooks(
    cleaned_paths: Iterable[Path],
    records: Iterable[AuditRecord],
    output_path: Path,
) -> int:
    """Combine cleaned data and audit details into one Excel workbook."""
    audit_records = list(records)
    combined_records: list[dict[str, Any]] = []
    offsets: dict[str, int] = {}

    for cleaned_path in cleaned_paths:
        offsets[str(cleaned_path)] = len(combined_records)
        offsets[cleaned_path.name] = len(combined_records)
        cleaned = load_workbook(cleaned_path, data_only=True, read_only=True)
        try:
            source_ws = (
                cleaned["cleaned_data"]
                if "cleaned_data" in cleaned.sheetnames
                else cleaned.worksheets[0]
            )
            rows = source_ws.iter_rows(values_only=True)
            headers = [clean_text(value) for value in next(rows, ())]
            for row in rows:
                record = {
                    header: value
                    for header, value in zip(headers, row)
                    if header
                }
                if any(not is_blank(value) for value in record.values()):
                    combined_records.append(record)
        finally:
            cleaned.close()

    output_columns = populated_columns(combined_records)
    workbook = Workbook()
    data_ws = workbook.active
    data_ws.title = "cleaned_data"
    data_ws.append(output_columns)
    for record in combined_records:
        data_ws.append([record.get(column) for column in output_columns])
    total_rows = len(combined_records)
    format_output(data_ws, total_rows, output_columns)

    audit_ws = workbook.create_sheet("cleaning_audit")
    audit_headers = [field.name for field in fields(AuditRecord)]
    audit_ws.append(audit_headers + ["consolidated_start_row", "consolidated_end_row"])
    for record in audit_records:
        values = asdict(record)
        offset = offsets.get(record.output_file)
        consolidated = [offset + row if offset is not None and row is not None else None
                        for row in (record.output_start_row, record.output_end_row)]
        audit_ws.append([json.dumps(values[key], ensure_ascii=False) if isinstance(values[key], (dict, list))
                         else values[key] for key in audit_headers] + consolidated)
    _format_audit_sheet(
        audit_ws,
        len(audit_records),
        {
            "input_file": 28,
            "sheet": 20,
            "header_row": 12,
            "status": 28,
            "score": 12,
            "data_rows": 12,
            "mapped_columns": 60,
            "unmapped_columns": 36,
            "output_file": 28,
            "block_id": 14,
            "source_range": 22,
            "warnings": 65,
            "excluded_columns": 36,
            "skipped_rows": 36,
            "merged_ranges": 30,
            "detection_method": 20,
        },
    )
    for column in ("G", "H"):
        for cell in audit_ws[column][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Review list only: these fields remain present in cleaned_data.
    unmapped_ws = workbook.create_sheet("unmapped_columns")
    unmapped_ws.append(
        ["input_file", "sheet", "block_id", "source_range", "header_row", "status", "unmapped_column"]
    )
    unmapped_rows = 0
    for record in audit_records:
        if record.unmapped_columns:
            for column in record.unmapped_columns:
                unmapped_ws.append(
                    [
                        record.input_file,
                        record.sheet,
                        record.block_id,
                        record.source_range,
                        record.header_row,
                        record.status,
                        column,
                    ]
                )
                unmapped_rows += 1
        elif record.status != "成功":
            unmapped_ws.append(
                [
                    record.input_file,
                    record.sheet,
                    record.block_id,
                    record.source_range,
                    record.header_row,
                    record.status,
                    "",
                ]
            )
            unmapped_rows += 1
    _format_audit_sheet(
        unmapped_ws,
        unmapped_rows,
        {
            "input_file": 28,
            "sheet": 20,
            "header_row": 12,
            "status": 28,
            "unmapped_column": 36,
        },
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return total_rows
