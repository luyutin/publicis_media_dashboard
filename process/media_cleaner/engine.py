"""Orchestrate merge expansion, region discovery, mapping and dated data export.

Only the output schema is shared with legacy media formatters.
"""

import argparse
import csv
import json
import sys
from dataclasses import asdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook

from .detection import discover_blocks
from .extraction import extract_block
from .grid import SheetGrid
from .mapping import map_blocks, read_dictionary
from .models import AuditRecord, Candidate, OllamaConfig
from .output import consolidate_cleaned_workbooks, format_output, populated_columns
from .workbooks import open_workbook
from .structure import refine_blocks


def list_workbook_sheets(workbook_data: bytes) -> list[str]:
    workbook = open_workbook(workbook_data, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def clean_workbook_sheets(
    input_path: Path, output_path: Path, aliases: dict[str, set[str]],
    scan_rows: int, sheet_names: Iterable[str], ollama: OllamaConfig | None = None,
    default_year: int | None = None,
) -> tuple[list[AuditRecord], int]:
    """Extract every dated region without mutating the input; 0 scans all rows."""
    if scan_rows < 0:
        raise ValueError("掃描列數不得小於 0；0 表示整張工作表")
    if default_year is not None and not 1900 <= default_year <= 2100:
        raise ValueError("補足日期的年份必須介於 1900–2100")
    if input_path.resolve() == output_path.resolve():
        raise ValueError("輸出路徑不得覆寫來源檔案")
    selected = list(dict.fromkeys(s for s in sheet_names if s))
    if not selected:
        raise ValueError("請至少選擇一個工作表。")
    workbook = open_workbook(input_path)
    audits, records = [], []
    try:
        missing = set(selected) - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"找不到工作表「{'、'.join(sorted(missing))}」。可用工作表：{'、'.join(workbook.sheetnames)}")
        for name in selected:
            try:
                grid = SheetGrid(workbook[name])
                blocks = discover_blocks(grid, scan_rows)
                blocks = refine_blocks(grid, blocks, aliases, ollama)
                map_blocks(grid, blocks, aliases, ollama)
                if not blocks:
                    audits.append(AuditRecord(str(input_path), name, None, "找不到資料表表頭", None, 0, {}, [], None))
                for block in blocks:
                    rows, audit = extract_block(grid, block, input_path, output_path, default_year)
                    if rows:
                        audit.output_start_row = len(records) + 2
                        records.extend(rows)
                        audit.output_end_row = len(records) + 1
                    audits.append(audit)
                if scan_rows and scan_rows < grid.max_row:
                    for audit in audits:
                        if audit.sheet == name:
                            audit.warnings.append(f"僅搜尋前 {scan_rows} 列表頭，下方可能還有未搜尋的區塊")
            except Exception as exc:
                audits.append(AuditRecord(str(input_path), name, None, f"處理失敗：{exc}", None, 0, {}, [], None))
    finally:
        workbook.close()
    if records:
        output = Workbook()
        try:
            sheet = output.active
            sheet.title = "cleaned_data"
            columns = populated_columns(records)
            sheet.append(columns)
            for record in records:
                sheet.append([record.get(column) for column in columns])
            format_output(sheet, len(records), columns)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output.save(output_path)
        finally:
            output.close()
    return audits, len(records)


def clean_workbook(input_path: Path, output_path: Path, aliases: dict[str, set[str]],
                   scan_rows: int = 0, sheet_name: str | None = None,
                   ollama: OllamaConfig | None = None, default_year: int | None = None):
    if sheet_name is None:
        workbook = open_workbook(input_path, read_only=True)
        try:
            sheet_name = workbook.active.title
        finally:
            workbook.close()
    return clean_workbook_sheets(input_path, output_path, aliases, scan_rows, [sheet_name], ollama, default_year)


def discover_inputs(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    return sorted(p for p in path.iterdir() if p.suffix.lower() in {".xls", ".xlsx"} and not p.name.startswith("~$")
                  and "template" not in p.stem.lower() and "cleaned" not in p.stem.lower())


def write_audit(records: list[AuditRecord], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "cleaning_audit.json").write_text(
        json.dumps([asdict(r) for r in records], ensure_ascii=False, indent=2), encoding="utf-8",
    )
    with (output_dir / "unmapped_columns.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["input_file", "sheet", "block_id", "source_range", "header_row", "status", "unmapped_column"])
        for record in records:
            columns = record.unmapped_columns or ([""] if record.status != "成功" else [])
            for column in columns:
                writer.writerow([record.input_file, record.sheet, record.block_id, record.source_range,
                                 record.header_row, record.status, column])


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="展開合併儲存格並擷取未知 Excel 中所有日期資料區塊")
    parser.add_argument("input", type=Path, help="單一 .xlsx / .xls 檔或資料夾")
    parser.add_argument("-o", "--output-dir", type=Path, default=Path("cleaned_output"))
    parser.add_argument("--dictionary", type=Path, default=Path("Report Template & All Format 字典.xlsx"))
    parser.add_argument("--scan-rows", type=int, default=0, help="表頭搜尋列數；0 表示整張工作表（預設）")
    parser.add_argument("--sheet", help="工作表名稱；預設為目前選取的工作表")
    parser.add_argument("--all-sheets", action="store_true", help="清理所有工作表")
    parser.add_argument("--default-year", type=int, help="僅補足缺少年份的日期，不改動完整日期")
    parser.add_argument("--ollama-model", default="qwen3.5:9b")
    parser.add_argument("--ollama-url", default="http://127.0.0.1:11434")
    parser.add_argument("--ollama-timeout", type=float, default=120)
    parser.add_argument("--no-ollama", action="store_true")
    args = parser.parse_args(argv)
    if args.sheet and args.all_sheets:
        parser.error("--sheet 與 --all-sheets 不可同時使用")
    return args


def main(argv=None) -> int:
    args = parse_args(argv)
    inputs = discover_inputs(args.input)
    if not inputs:
        print(f"找不到可處理的 .xlsx：{args.input}", file=sys.stderr)
        return 2
    aliases = read_dictionary(args.dictionary)
    config = OllamaConfig(model=args.ollama_model, url=args.ollama_url,
                          timeout=args.ollama_timeout, enabled=not args.no_ollama,
                          force_all=not args.no_ollama)
    audits, outputs = [], []
    for path in inputs:
        output = args.output_dir / f"{path.stem}_cleaned.xlsx"
        try:
            if args.all_sheets:
                sheets = list_workbook_sheets(path.read_bytes())
                records, count = clean_workbook_sheets(path, output, aliases, args.scan_rows, sheets, config, args.default_year)
            else:
                records, count = clean_workbook(path, output, aliases, args.scan_rows, args.sheet, config, args.default_year)
        except (OSError, ValueError) as exc:
            records = [AuditRecord(str(path), args.sheet or "", None, f"處理失敗：{exc}", None, 0, {}, [], None)]
            count = 0
        audits.extend(records)
        if count:
            outputs.append(output)
        print(f"{path.name}: {count} 列，{len(records)} 筆區塊稽核")
    write_audit(audits, args.output_dir)
    consolidate_cleaned_workbooks(outputs, audits, args.output_dir / "cleaned_media_results.xlsx")
    return 0 if outputs else 1


if __name__ == "__main__":
    raise SystemExit(main())
