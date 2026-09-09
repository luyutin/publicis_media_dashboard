"""Validate and extract dated records strictly inside each discovered region."""

import re
from collections import Counter

from openpyxl.utils import get_column_letter

from .grid import clean_text, normalize, is_blank, is_error, is_summary, date_parts, coerce_excel_date
from .mapping import excluded_header
from .models import AuditRecord
from .settings import SOURCE_COLUMN, TARGET_COLUMNS


def _context_fields(grid, block):
    labels = {
        "刊登平台": "Platform", "publisherplatform": "Platform",
        "刊登版位": "Placement", "廣告版位": "Placement", "placement": "Placement",
        "客戶案名": "Campaign name", "campaignname": "Campaign name",
    }
    fields = {}
    if block.group_label and not is_summary(block.group_label):
        fields["Placement"] = block.group_label
    for row in grid.context(block.header_row, block.start_column, block.end_column):
        for index, value in enumerate(row):
            target = labels.get(normalize(clean_text(value).rstrip(":：")))
            if target:
                following = next((v for v in row[index + 1:] if not is_blank(v) and v != value), None)
                if following is not None:
                    fields[target] = following
    return fields


def _resolve_year(grid, block, date_column, default_year):
    if default_year is not None:
        return default_year, "使用指定年份補足缺少年份的日期"
    years = set()
    for row in range(block.header_row + 1, block.end_row + 1):
        parts = date_parts(grid.get(row, date_column))
        if parts and parts[0]:
            years.add(parts[0])
    if len(years) == 1:
        return years.pop(), "使用同區塊完整日期的年份補足日期"
    if len(years) > 1:
        return None, None
    # Only explicitly labelled years in local metadata; never use the current year.
    for row in reversed(grid.context(block.header_row, block.start_column, block.end_column)):
        text = " ".join(clean_text(v) for v in row)
        labelled = set(re.findall(r"(?<!\d)((?:19|20)\d{2})\s*年", text))
        labelled.update(re.findall(r"(?:year|年度|年份)\s*[:：]?\s*((?:19|20)\d{2})(?!\d)", text, re.I))
        if len(labelled) == 1:
            return int(labelled.pop()), "使用區塊上方標示的年份補足日期"
        if len(labelled) > 1:
            return None, None
    return None, None


def extract_block(grid, block, input_path, output_path, default_year=None):
    area = f"{get_column_letter(block.start_column)}{block.header_start}:{get_column_letter(block.end_column)}{block.end_row}"
    record = AuditRecord(
        input_file=str(input_path), sheet=grid.sheet, header_row=block.header_row,
        status="成功", score=block.score, data_rows=0,
        mapped_columns={f"{get_column_letter(c)}: {block.headers[c - 1]}": target for c, target in block.mapped.items()},
        unmapped_columns=block.unmapped, output_file=None, block_id=block.block_id,
        source_range=area, detection_method=block.method,
        excluded_columns=block.excluded_columns, warnings=list(block.warnings),
        merged_ranges=[str(m) for m in grid.merges if m.min_col <= block.end_column
                       and m.max_col >= block.start_column and m.min_row <= block.end_row
                       and m.max_row >= block.header_start],
    )
    date_column = next((c for c, target in block.mapped.items() if target == "Date"), None)
    if date_column is not None:
        record.date_source_range = f"{get_column_letter(date_column)}{block.header_row + 1}:{get_column_letter(date_column)}{block.end_row}"
    if block.group_label and is_summary(block.group_label):
        record.status = "排除：彙總欄群組"
        return [], record
    if date_column is None:
        record.status = "排除：沒有可辨識的日期欄位"
        return [], record
    output_fields = _output_fields(grid, block, date_column, record)
    mapped_data_fields = [target for target in block.mapped.values() if target != "Date"]
    if not mapped_data_fields and not output_fields:
        record.status = "排除：日期以外沒有可保留的欄位"
        return [], record
    year, year_note = _resolve_year(grid, block, date_column, default_year)
    context = _context_fields(grid, block)
    counts = Counter()
    rows = []
    summary_columns = [c for c in range(block.start_column, block.end_column + 1)
                       if not excluded_header(block.headers[c - 1])]
    if date_column not in summary_columns:
        summary_columns.append(date_column)
    for row in range(block.header_row + 1, block.end_row + 1):
        values = [grid.get(row, c) for c in summary_columns]
        if all(is_blank(v) for v in values):
            counts["空白列"] += 1
            continue
        if any(is_summary(v) for v in values):
            counts["彙總或目標列"] += 1
            continue
        raw_date = grid.get(row, date_column)
        parsed = coerce_excel_date(raw_date, year=year, epoch=grid.epoch)
        if parsed is None:
            parts = date_parts(raw_date)
            reason = "日期缺少年份" if parts and parts[0] is None and year is None else "無效或缺少日期"
            counts[reason] += 1
            continue
        parts = date_parts(raw_date)
        if parts and parts[0] is None and year_note and year_note not in record.warnings:
            record.warnings.append(year_note)
        # Only declare fields that came from this source block (plus inferred
        # context). This lets empty source columns survive schema collection.
        item = {SOURCE_COLUMN: input_path.name}
        item.update(context)
        for col, target in block.mapped.items():
            value = grid.get(row, col)
            item[target] = None if is_blank(value) or is_error(value) else value
        for col, output_name in output_fields.items():
            value = grid.get(row, col)
            item[output_name] = None if is_blank(value) or is_error(value) else value
        item["Date"] = parsed
        data_fields = mapped_data_fields + list(output_fields.values())
        if not any(not is_blank(item.get(field)) for field in data_fields):
            counts["日期以外沒有有效資料"] += 1
            continue
        rows.append(item)
    record.skipped_rows = dict(counts)
    if counts["日期缺少年份"]:
        record.warnings.append("部分日期缺少年份且無法從區塊確認；請在進階設定指定年份後重跑")
    record.data_rows = len(rows)
    record.output_file = str(output_path) if rows else None
    if not rows:
        record.status = "排除：沒有有效日期資料列"
    return rows, record


def _output_fields(grid, block, date_column, audit):
    """Name every unmapped source column without overwriting another field."""
    result = {}
    used = set(TARGET_COLUMNS) | {SOURCE_COLUMN}
    for col in range(block.start_column, block.end_column + 1):
        if col == date_column or col in block.mapped:
            continue
        header = block.headers[col - 1]
        if excluded_header(header):
            continue
        # A horizontal merge repeats one physical source cell across columns.
        # Keep it once; vertical merges remain filled on every data row.
        duplicate_merge = next((
            previous for previous in range(block.start_column, col)
            if (previous == date_column or previous in block.mapped or previous in result)
            and _same_merged_source(grid, block, previous, col)
        ), None)
        if duplicate_merge is not None:
            message = f"{get_column_letter(col)}: {header}（與 {get_column_letter(duplicate_merge)} 為同一合併儲存格）"
            if message not in audit.excluded_columns:
                audit.excluded_columns.append(message)
            continue
        original = clean_text(grid.get(block.header_row, col)) or clean_text(header)
        if not original:
            original = f"未命名欄_{get_column_letter(col)}"
        output_name = original
        if output_name in used:
            output_name = f"{original} [{get_column_letter(col)}]"
            audit.warnings.append(f"來源欄名「{original}」重複或與保留欄位同名，輸出為「{output_name}」")
        used.add(output_name)
        result[col] = output_name
    return result


def _same_merged_source(grid, block, left, right):
    """True only when all populated values come from the same merged cells."""
    found = False
    for row in range(block.header_row + 1, block.end_row + 1):
        left_value, right_value = grid.get(row, left), grid.get(row, right)
        if is_blank(left_value) and is_blank(right_value):
            continue
        left_origin = grid.origins.get((row, left))
        right_origin = grid.origins.get((row, right))
        if left_origin is None or left_origin != right_origin:
            return False
        found = True
    return found
