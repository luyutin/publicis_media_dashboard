"""Read XLSX and legacy XLS into the same value/merge workbook interface."""

import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH


def open_workbook(source: Path | bytes, read_only: bool = False):
    data = source if isinstance(source, bytes) else source.read_bytes()
    if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return load_workbook(io.BytesIO(data), data_only=True, read_only=read_only)
    import xlrd

    legacy = xlrd.open_workbook(file_contents=data, formatting_info=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.epoch = MAC_EPOCH if legacy.datemode else WINDOWS_EPOCH
    try:
        for source_sheet in legacy.sheets():
            sheet = workbook.create_sheet(source_sheet.name)
            for row in range(source_sheet.nrows):
                for col in range(source_sheet.ncols):
                    cell = source_sheet.cell(row, col)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, legacy.datemode)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#VALUE!")
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    sheet.cell(row + 1, col + 1, value)
            for first_row, last_row, first_col, last_col in source_sheet.merged_cells:
                sheet.merge_cells(start_row=first_row + 1, end_row=last_row,
                                  start_column=first_col + 1, end_column=last_col)
        return workbook
    finally:
        legacy.release_resources()
