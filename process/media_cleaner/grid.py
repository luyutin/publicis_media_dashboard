"""Read-only logical worksheet: expand actual merges, never forward-fill blanks."""

import math
import re
from datetime import date, datetime
from typing import Any

from openpyxl.utils.datetime import from_excel, WINDOWS_EPOCH


def clean_text(value: Any) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def normalize(value: Any) -> str:
    text = clean_text(value).casefold().replace("（", "(").replace("）", ")")
    return re.sub(r"[\s_:/：,，。·・\-]+", "", text)


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) or clean_text(value) in {"", "--", "—"}


def is_error(value: Any) -> bool:
    return isinstance(value, str) and value in {
        "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!",
    }


SUMMARY_LABELS = {
    "total", "grandtotal", "overalltotal", "subtotal", "totalgeneral",
    "合計", "總計", "小計", "總數", "加總", "總和", "總合", "总计", "合计", "总数",
    "委刊目標", "與目標差距", "預估成效", "達成率", "estimate",
}


def is_summary(value: Any) -> bool:
    # Exact labels avoid dropping real metrics such as Total PPL / total users.
    return normalize(clean_text(value).strip(" :：-–—_")) in SUMMARY_LABELS


def date_parts(value: Any) -> tuple[int | None, int, int] | None:
    if isinstance(value, (date, datetime)):
        return value.year, value.month, value.day
    if not isinstance(value, str):
        return None
    text = value.strip()
    full = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]00:00:00)?", text)
    western = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    partial = re.fullmatch(r"(\d{1,2})[/月](\d{1,2})日?", text)
    if full:
        parts = tuple(map(int, full.groups()))
    elif western:
        month, day, year = map(int, western.groups())
        parts = year, month, day
    elif partial:
        month, day = map(int, partial.groups())
        parts = None, month, day
    else:
        return None
    try:
        date(parts[0] or 2000, parts[1], parts[2])
    except ValueError:
        return None
    return parts


def coerce_excel_date(value: Any, year: int | None = None, epoch=WINDOWS_EPOCH,
                      allow_serial: bool = True) -> date | None:
    parts = date_parts(value)
    if parts:
        resolved_year = parts[0] or year
        if resolved_year is None:
            return None
        try:
            return date(resolved_year, parts[1], parts[2])
        except ValueError:
            return None
    if allow_serial and isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 <= value <= 80000:
        return from_excel(value, epoch=epoch).date()
    return None


def numeric(value: Any) -> bool:
    if isinstance(value, bool) or is_blank(value):
        return False
    if isinstance(value, (int, float)):
        return True
    return bool(re.fullmatch(r"[-+]?[$＄€£]?\d[\d,]*(?:\.\d+)?%?", clean_text(value)))


class SheetGrid:
    def __init__(self, worksheet):
        self.sheet = worksheet.title
        self.epoch = worksheet.parent.epoch
        self.values = {
            (cell.row, cell.column): cell.value
            for row in worksheet.iter_rows() for cell in row if not is_blank(cell.value)
        }
        self.merges = list(worksheet.merged_cells.ranges)
        self.origins: dict[tuple[int, int], tuple[int, int]] = {}
        for area in self.merges:
            anchor = (area.min_row, area.min_col)
            value = self.values.get(anchor)
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    self.origins[row, col] = anchor
                    if not is_blank(value):
                        self.values[row, col] = value
        self.max_row = max((r for r, _ in self.values), default=0)
        self.max_column = max((c for _, c in self.values), default=0)

    def get(self, row: int, col: int):
        return self.values.get((row, col))

    def row(self, row: int, start: int, end: int) -> list[Any]:
        return [self.get(row, col) for col in range(start, end + 1)]

    def header(self, row: int, col: int) -> tuple[str, int]:
        """Include horizontal merged parents in a multi-level header."""
        parts = [clean_text(self.get(row, col))]
        first = self.origins.get((row, col), (row, col))[0]
        for above in range(row - 1, max(0, row - 3), -1):
            area = next((m for m in self.merges if m.min_row <= above <= m.max_row
                         and m.min_col <= col <= m.max_col and m.max_col > m.min_col), None)
            if not area:
                break
            parent = clean_text(self.get(above, col))
            if parent and parent not in parts:
                parts.insert(0, parent)
                first = min(first, area.min_row)
        return " | ".join(part for part in parts if part), first

    def context(self, header_row: int, start: int, end: int) -> list[list[Any]]:
        # Keep context local in BOTH axes, so adjacent placements cannot leak.
        return [self.row(r, start, end) for r in range(max(1, header_row - 12), header_row)]
