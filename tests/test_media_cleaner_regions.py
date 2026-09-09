"""Behavioural regressions for irregular media layouts, using real XLSX files."""

import json
import tempfile
import unittest
from contextlib import nullcontext
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH, to_excel

from process import media_cleaner
from process.media_cleaner.detection import discover_blocks
from process.media_cleaner.grid import SheetGrid
from process.media_cleaner.mapping import map_blocks
from process.media_cleaner.structure import refine_blocks


class RegionCleanerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.book = Workbook()
        self.addCleanup(self.book.close)
        self.ws = self.book.active
        self.ws.title = "Report"

    def run_cleaner(self, year=None, llm=False, scan=0, layout=False):
        source = self.root / "source.xlsx"
        output = self.root / "result.xlsx"
        self.book.save(source)
        before = source.read_bytes()
        with (nullcontext() if layout else patch("process.media_cleaner.engine.refine_blocks", side_effect=lambda grid, seeds, aliases, config: seeds)):
            audit, count = media_cleaner.clean_workbook_sheets(
                source, output, media_cleaner.read_dictionary(None), scan, ["Report"],
                media_cleaner.OllamaConfig(enabled=llm), default_year=year,
            )
        self.assertEqual(source.read_bytes(), before, "Never edit source workbook")
        failures = [r.status for r in audit if r.status.startswith("處理失敗")]
        self.assertEqual(failures, [])
        records = []
        if count:
            result = load_workbook(output, data_only=True)
            try:
                values = list(result.active.values)
                records = [dict(zip(values[0], row)) for row in values[1:]]
            finally:
                result.close()
        self.assertEqual(count, len(records))
        self.assertTrue(all(isinstance(row.get("Date"), datetime) for row in records))
        return audit, records

    def put(self, row, col, values):
        for offset, value in enumerate(values):
            self.ws.cell(row, col + offset, value)

    def test_four_parallel_tables_have_independent_totals_and_context(self):
        for col, platform, placement in ((1, "手機網", "300x250"), (6, "手機網", "300x250"),
                                         (11, "APP", "bottom"), (16, "手機網", "開機蓋板")):
            self.put(1, col, ["刊登平台", platform])
            self.ws.merge_cells(start_row=1, start_column=col+1, end_row=1, end_column=col+3)
            self.put(2, col, ["刊登版位", placement])
            self.put(5, col, ["日期 項目", "曝光", "點擊數", "點擊率"])
            self.put(6, col, ["2024/2/1", col * 100, 10, 0.1])
            self.put(7, col, ["2024/2/2", col * 200, 20, 0.1])
            self.put(8, col, ["總數" if col == 16 else "2024/2/3", col * 300, 30, 0.1])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 11)
        self.assertEqual([a.data_rows for a in audit], [3, 3, 3, 2])
        self.assertEqual({r["Platform"] for r in records}, {"手機網", "APP"})
        self.assertEqual([r["Placement"] for r in records[-2:]], ["開機蓋板"] * 2)
        self.assertEqual(audit[-1].skipped_rows["彙總或目標列"], 1)
        self.assertEqual(audit[0].output_start_row, 2)
        self.assertEqual(audit[-1].output_end_row, 12)

    def test_stacked_demographics_excluded_and_lower_daily_table_found(self):
        for row, label, value in ((1, "Date", "2025/7/28"), (5, "年齡", "18–24"),
                                  (9, "性別", "女"), (150, "Date", "2025/8/1")):
            self.put(row, 1, [label, "Impressions", "Clicks"])
            self.put(row + 1, 1, [value, row * 100, 10])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 2)
        self.assertEqual(len(audit), 4)
        self.assertEqual([a.status for a in audit[1:3]], ["排除：沒有可辨識的日期欄位"] * 2)
        self.assertEqual(records[-1]["Date"].date(), date(2025, 8, 1))

    def test_adjacent_tables_without_spacer_and_offset_headers(self):
        self.put(1, 1, ["Date", "Impressions", "Date", "Clicks"])
        self.put(2, 1, ["2025/1/1", 100, "2025/1/2", 20])
        self.put(3, 1, ["2025/1/3", 200, "Grand Total", 20])
        self.put(8, 5, ["Date", "Views"])
        self.put(9, 5, ["2025/1/4", 50])
        audit, records = self.run_cleaner()
        self.assertEqual([a.data_rows for a in audit], [2, 1, 1])
        self.assertEqual(sum(r.get("Impressions") or 0 for r in records), 300)

    def test_merged_dimensions_are_filled_but_unmerged_blanks_are_not(self):
        self.put(1, 1, ["Daily Performance", "Campaign Objective", "Spend", "Result type",
                        "Result", "Cost per result type", "Cost per result", "Impressions", "Campaign name"])
        for row in range(2, 12):
            self.put(row, 1, [date(2025, 7, 23 + row) if row < 9 else date(2025, 8, row - 8),
                              None, 500, None, 540000, None, 1.2, 540000, "A" if row == 2 else None])
        for col, value in ((2, "Awareness"), (4, "Impressions"), (6, "CPM")):
            self.ws.cell(2, col, value)
            self.ws.merge_cells(start_row=2, start_column=col, end_row=11, end_column=col)
        grid = SheetGrid(self.ws)
        self.assertEqual(grid.get(11, 2), "Awareness")
        self.assertEqual(grid.get(11, 6), "CPM")
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 10)
        self.assertTrue(all(r["Campaign Objective"] == "Awareness" for r in records))
        self.assertEqual(records[0]["Campaign name"], "A")
        self.assertIsNone(records[1]["Campaign name"])
        self.assertTrue(all(r["Spent (TWD)"] == 500 for r in records))
        self.assertIn("B2:B11", audit[0].merged_ranges)
        self.assertIn("Cost per result", audit[0].unmapped_columns)

    def test_multilevel_merged_header_and_total_column(self):
        self.put(1, 1, ["Date", "Performance", None, "Total"])
        self.ws.merge_cells("A1:A2")
        self.ws.merge_cells("B1:C1")
        self.ws.merge_cells("D1:D2")
        self.put(2, 2, ["Impressions", "Clicks"])
        self.put(3, 1, ["2025/7/28", 100, 20, 120])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Impressions"], 100)
        self.assertEqual(records[0]["Clicks (all)"], 20)
        self.assertEqual(audit[0].source_range, "A1:D3")
        self.assertTrue(audit[0].excluded_columns)

    def test_merged_total_parent_excludes_all_children(self):
        self.put(1, 1, ["Date", "Impressions", "Total"])
        self.ws.merge_cells("A1:A2")
        self.ws.merge_cells("B1:B2")
        self.ws.merge_cells("C1:D1")
        self.put(2, 3, ["Clicks", "Views"])
        self.put(3, 1, ["2025/1/1", 100, 20, 30])
        audit, records = self.run_cleaner()
        self.assertEqual(set(records[0]), {"Date", "Impressions", "Source"})
        self.assertEqual(len(audit[0].excluded_columns), 2)

    def test_total_metric_names_are_not_aggregate_columns(self):
        self.ws.append(["Date", "Total PPL", "total users", "Total"])
        self.ws.append(["2025/1/1", 100, 50, "Total"])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Total PPL"], 100)
        self.assertEqual(records[0]["User number"], 50)

    def test_unknown_year_excluded_and_explicit_year_only_fills_missing(self):
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["7月28日", 100])
        audit, records = self.run_cleaner()
        self.assertEqual(records, [])
        self.assertEqual(audit[0].skipped_rows["日期缺少年份"], 1)
        self.ws.append(["2024/7/29", 200])
        audit, records = self.run_cleaner(year=2025)
        self.assertEqual([r["Date"].year for r in records], [2025, 2024])

    def test_local_year_metadata_and_date_serial_epoch(self):
        self.ws.append(["年度：2025"])
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["7月28日", 100])
        _, records = self.run_cleaner()
        self.assertEqual(records[0]["Date"].year, 2025)
        self.book.epoch = MAC_EPOCH
        self.ws.cell(3, 1, to_excel(datetime(2025, 7, 28), epoch=MAC_EPOCH))
        _, records = self.run_cleaner()
        self.assertEqual(records[0]["Date"].date(), date(2025, 7, 28))

    def test_invalid_dates_and_excel_errors_do_not_become_records(self):
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["2025/2/30", 100])
        self.ws.append(["2025/2/28", "#DIV/0!"])
        self.ws.append(["2025/3/1", 0])
        _, records = self.run_cleaner()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Impressions"], 0)

    def test_llm_maps_unseen_metric_without_alias_and_rejects_contradictions(self):
        self.ws.append(["出街日", "觸達展示總筆數", "6s Views", "Cost per result", "Total"])
        self.ws.append(["2025/1/1", 123, 20, 2.5, 145])
        response = {"blocks": [{"id": "table_1", "mappings": [
            {"column": 2, "target": "Impressions", "confidence": 0.95},
            {"column": 3, "target": '3" Video Views', "confidence": 0.99},
            {"column": 4, "target": "Spent (TWD)", "confidence": 0.99},
            {"column": 5, "target": "Clicks (all)", "confidence": 0.99},
            {"column": 100, "target": "Reach", "confidence": 1},
        ]}]}
        with patch("process.media_cleaner.mapping.call_ollama_json", return_value=response):
            audit, records = self.run_cleaner(llm=True)
        self.assertEqual(set(records[0]), {"Date", "Impressions", "Source", "6s Views", "Cost per result"})
        self.assertEqual(records[0]["Impressions"], 123)
        self.assertEqual(records[0]["6s Views"], 20)
        self.assertEqual(records[0]["Cost per result"], 2.5)
        self.assertEqual(audit[0].detection_method, "rules+llm")

    def test_llm_failure_keeps_rules_and_surfaces_warning(self):
        self.ws.append(["Date", "Impressions", "備註"])
        self.ws.append(["2025/1/1", 100, "sample"])
        with patch("process.media_cleaner.mapping.call_ollama_json", side_effect=RuntimeError("offline")):
            audit, records = self.run_cleaner(llm=True)
        self.assertEqual(len(records), 1)
        self.assertTrue(any("offline" in warning for warning in audit[0].warnings))

    def test_force_all_sends_fully_alias_mapped_dated_block_to_llm(self):
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["2025/1/1", 100])
        grid = SheetGrid(self.ws)
        blocks = discover_blocks(grid)
        config = media_cleaner.OllamaConfig(enabled=True, force_all=True)
        response = {"blocks": [{"id": blocks[0].block_id, "mappings": []}]}
        with patch("process.media_cleaner.mapping.call_ollama_json", return_value=response) as mocked:
            map_blocks(grid, blocks, media_cleaner.read_dictionary(None), config)
        self.assertEqual(mocked.call_count, 1)

    def test_force_all_also_sends_block_without_date_before_excluding_it(self):
        self.ws.append(["年齡", "Impressions"])
        self.ws.append(["18-24", 100])
        grid = SheetGrid(self.ws)
        blocks = discover_blocks(grid)
        config = media_cleaner.OllamaConfig(enabled=True, force_all=True)
        response = {"blocks": [{"id": blocks[0].block_id, "mappings": []}]}
        with patch("process.media_cleaner.mapping.call_ollama_json", return_value=response) as mocked:
            map_blocks(grid, blocks, media_cleaner.read_dictionary(None), config)
        self.assertEqual(mocked.call_count, 1)
        self.assertEqual(blocks[0].method, "rules+llm")

    def test_mapping_failure_does_not_disable_later_batches(self):
        for index in range(7):
            row = 1 + index * 4
            self.put(row, 1, ["Date", "陌生展示量"])
            self.put(row + 1, 1, ["2025/1/1", index + 1])
        grid = SheetGrid(self.ws)
        blocks = discover_blocks(grid)
        config = media_cleaner.OllamaConfig(enabled=True)

        def answer(prompt, _config):
            if answer.calls == 0:
                answer.calls += 1
                raise RuntimeError("first batch offline")
            answer.calls += 1
            ids = [block.block_id for block in blocks[6:]]
            return {"blocks": [{"id": block_id, "mappings": [
                {"column": 2, "target": "Impressions", "confidence": 0.9}
            ]} for block_id in ids]}

        answer.calls = 0
        with patch("process.media_cleaner.mapping.call_ollama_json", side_effect=answer) as mocked:
            map_blocks(grid, blocks, media_cleaner.read_dictionary(None), config)

        self.assertEqual(mocked.call_count, 2)
        self.assertTrue(config.enabled)
        self.assertTrue(all("first batch offline" in block.warnings[0] for block in blocks[:6]))
        self.assertEqual(blocks[6].mapped[2], "Impressions")
        self.assertFalse(blocks[6].warnings)

    def test_structure_failure_does_not_disable_later_header_rows(self):
        self.put(1, 1, ["Date", "陌生指標"])
        self.put(2, 1, ["2025/1/1", 10])
        self.put(5, 1, ["Date", "另一指標"])
        self.put(6, 1, ["2025/1/2", 20])
        grid = SheetGrid(self.ws)
        seeds = discover_blocks(grid)
        config = media_cleaner.OllamaConfig(enabled=True)
        valid = {"table_count": 1, "blocks": [{
            "start_column": 1, "end_column": 2, "date_column": 1,
            "label_row": 0, "label_column": 0, "confidence": 0.95,
        }]}
        with patch(
            "process.media_cleaner.structure.call_ollama_layout",
            side_effect=[RuntimeError("first row offline"), valid],
        ) as mocked:
            blocks = refine_blocks(grid, seeds, media_cleaner.read_dictionary(None), config)

        self.assertEqual(mocked.call_count, 2)
        self.assertTrue(config.enabled)
        self.assertTrue(any("first row offline" in warning for warning in blocks[0].warnings))
        self.assertEqual(blocks[1].method, "llm-layout")

    def test_llm_malformed_conflicts_and_low_confidence_do_not_corrupt_rules(self):
        self.ws.append(["Date", "Impressions", "甲", "乙", "丙"])
        self.ws.append(["2025/1/1", 100, 20, 30, 40])
        response = {"blocks": [{"id": "table_1", "mappings": [
            None, {"column": [], "target": "Clicks (all)", "confidence": 1},
            {"column": 3, "target": "Clicks (all)", "confidence": 0.9},
            {"column": 4, "target": "Clicks (all)", "confidence": 0.9},
            {"column": 5, "target": "Reach", "confidence": 0.4},
            {"column": 2, "target": "Spent (TWD)", "confidence": 0.99},
        ]}]}
        with patch("process.media_cleaner.mapping.call_ollama_json", return_value=response):
            audit, records = self.run_cleaner(llm=True)
        self.assertEqual(set(records[0]), {"Date", "Impressions", "Source", "甲", "乙", "丙"})
        self.assertTrue(audit[0].warnings)

    def test_every_unknown_block_reaches_llm_not_just_first_batch(self):
        for index in range(14):
            self.put(1 + index * 4, 1, ["Date", "未知展示數"])
            self.put(2 + index * 4, 1, ["2025/1/1", index])
        def answer(prompt, config):
            import re
            ids = re.findall(r'"id": "(table_\d+)"', prompt)
            return {"blocks": [{"id": i, "mappings": [
                {"column": 2, "target": "Impressions", "confidence": 0.9}
            ]} for i in ids]}
        with patch("process.media_cleaner.mapping.call_ollama_json", side_effect=answer) as mocked:
            audit, records = self.run_cleaner(llm=True)
        self.assertEqual(len(records), 14)
        self.assertEqual(mocked.call_count, 3)

    def test_block_audits_survive_json_csv_and_consolidated_excel(self):
        self.ws.append(["Date", "Impressions", "Total"])
        self.ws.append(["2025/1/1", 100, 100])
        self.ws.append(["總計", 100, 100])
        audit, records = self.run_cleaner()
        media_cleaner.write_audit(audit, self.root)
        exported = json.loads((self.root / "cleaning_audit.json").read_text())
        self.assertEqual(exported[0]["source_range"], "A1:C3")
        self.assertEqual(exported[0]["skipped_rows"]["彙總或目標列"], 1)
        result = self.root / "combined.xlsx"
        media_cleaner.consolidate_cleaned_workbooks([self.root / "result.xlsx"], audit, result)
        book = load_workbook(result, data_only=True)
        try:
            values = list(book["cleaning_audit"].values)
            record = dict(zip(values[0], values[1]))
            self.assertEqual(record["block_id"], "table_1")
            self.assertEqual(record["consolidated_start_row"], 2)
            self.assertEqual(record["consolidated_end_row"], 2)
        finally:
            book.close()

    def test_header_search_limit_warns_but_does_not_truncate_data(self):
        self.ws.append(["Date", "Impressions"])
        for day in range(1, 10):
            self.ws.append([f"2025/1/{day}", day])
        audit, records = self.run_cleaner(scan=2)
        self.assertEqual(len(records), 9)
        self.assertTrue(any("僅搜尋" in w for w in audit[0].warnings))

    def test_text_heavy_data_rows_are_never_headers(self):
        self.ws.append(["Date", "Campaign name", "Audience", "Platform", "Placement", "Product", "Channel", "Media", "Impressions"])
        for day in range(1, 5):
            self.ws.append([f"2025/1/{day}", "Summer", "Adults", "APP", "bottom", "Tea", "Social", "Meta", 100])
        audit, records = self.run_cleaner()
        self.assertEqual(len(audit), 1)
        self.assertEqual(len(records), 4)

    def test_horizontal_merges_fill_every_cell_without_duplicating_metrics(self):
        self.put(1, 1, ["Date", None, "Impressions", None, "Clicks"])
        self.ws.merge_cells("A1:B1")
        self.ws.merge_cells("C1:D1")
        for row in (2, 3):
            self.put(row, 1, [f"2025/1/{row}", None, 100 * row, None, 20])
            self.ws.merge_cells(f"A{row}:B{row}")
            self.ws.merge_cells(f"C{row}:D{row}")
        grid = SheetGrid(self.ws)
        self.assertEqual(grid.get(3, 4), 300)
        audit, records = self.run_cleaner()
        self.assertEqual(len(audit), 1)
        self.assertEqual(len(records), 2)
        self.assertEqual([r["Impressions"] for r in records], [200, 300])
        self.assertNotIn("Impressions [D]", records[0])

    def test_llm_cannot_turn_numeric_demographics_into_excel_dates(self):
        self.ws.append(["分組代碼", "Impressions", "備註"])
        self.ws.append([45000, 100, "x"])
        response = {"blocks": [{"id": "table_1", "mappings": [
            {"column": 1, "target": "Date", "confidence": 0.99},
        ]}]}
        with patch("process.media_cleaner.mapping.call_ollama_json", return_value=response):
            audit, records = self.run_cleaner(llm=True)
        self.assertEqual(records, [])
        self.assertEqual(audit[0].status, "排除：沒有可辨識的日期欄位")

    def test_merged_year_title_can_supply_year(self):
        self.ws.append(["2025年成效報告"])
        self.ws.merge_cells("A1:B1")
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["7月28日", 100])
        _, records = self.run_cleaner()
        self.assertEqual(records[0]["Date"].year, 2025)

    def test_numeric_serial_date_with_many_dimensions_is_not_a_header(self):
        self.ws.append(["Date", "Campaign name", "Audience", "Platform", "Placement", "Product", "Channel", "Media", "Impressions"])
        for day in range(1, 5):
            self.ws.append([to_excel(datetime(2025, 1, day)), "Summer", "Adults", "APP", "bottom", "Tea", "Social", "Meta", 100])
        audit, records = self.run_cleaner()
        self.assertEqual(len(audit), 1)
        self.assertEqual(len(records), 4)

    def test_columns_beyond_old_100_column_limit_are_processed(self):
        self.put(1, 120, ["Date", "Impressions"])
        self.put(2, 120, ["2025/1/1", 100])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 1)
        self.assertEqual(audit[0].source_range, "DP1:DQ2")

    def test_shared_date_groups_exclude_total_and_keep_late_starting_placements(self):
        self.put(1, 1, ["版位", "合計", None, None, "APP", None, None, "Web"])
        for area in ("B1:D1", "E1:G1", "H1:J1"):
            self.ws.merge_cells(area)
        self.put(4, 1, ["日期/成效", "Impressions", "Clicks", "CTR", "Impressions", "Clicks", "CTR", "Impressions", "Clicks", "CTR"])
        for row in range(5, 9):
            self.put(row, 1, [date(2025, 1, row), 300, 30, 0.1, 100, 10, 0.1])
        self.put(8, 8, [200, 20, 0.1])
        self.put(9, 1, ["合計", 1200, 120, 0.1, 400, 40, 0.1, 200, 20, 0.1])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 5)
        self.assertEqual(sum(r["Impressions"] for r in records), 600)
        self.assertEqual([r["Placement"] for r in records], ["APP"] * 4 + ["Web"])
        self.assertEqual(audit[0].status, "排除：彙總欄群組")
        self.assertEqual(audit[-1].date_source_range, "A5:A9")

    def test_shared_axis_with_estimate_row_and_schedule_above_metrics(self):
        self.put(1, 1, ["Position", "總合", None, None, "APP", None, None, "Web"])
        for area in ("B1:D1", "E1:G1", "H1:J1"):
            self.ws.merge_cells(area)
        self.put(2, 1, ["Date", None, None, None, "2/14-2/27", None, None, "2/20-2/27"])
        self.ws.merge_cells("E2:G2")
        self.ws.merge_cells("H2:J2")
        self.put(4, 1, ["Estimate", "Imps", "Clicks", "CTR", "Imps", "Clicks", "CTR", "Imps", "Clicks", "CTR"])
        self.put(5, 1, ["Estimate", 300, 30, 0.1, 100, 10, 0.1, 200, 20, 0.1])
        self.put(6, 1, [date(2025, 2, 20), 300, 30, 0.1, 100, 10, 0.1, 200, 20, 0.1])
        _, records = self.run_cleaner()
        self.assertEqual(len(records), 2)
        self.assertEqual({r["Placement"] for r in records}, {"APP", "Web"})

    def test_explicit_twd_spend_wins_over_unspecified_spend(self):
        self.ws.append(["Date", "Impression", "Spend", "Spend\n($NTD)"])
        self.ws.append(["2025/7/28", 100, 10, 310])
        audit, records = self.run_cleaner()
        self.assertEqual(records[0]["Spent (TWD)"], 310)
        self.assertEqual(records[0]["Spend"], 10)
        self.assertTrue(audit[0].warnings)

    def test_internal_blank_rows_do_not_truncate_daily_data(self):
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["2025/1/1", 100])
        self.ws.append([])
        self.ws.append([])
        self.ws.append(["2025/1/2", 200])
        _, records = self.run_cleaner()
        self.assertEqual(len(records), 2)

    def test_llm_structure_handles_date_anywhere_and_different_metric_orders(self):
        # Expected layouts are provided by a stub: this verifies the parser's
        # generality, not the real model's ability to discover those layouts.
        for date_position in (0, 2, 4):
            with self.subTest(date_position=date_position):
                self.ws.delete_rows(1, self.ws.max_row)
                headers = ["Impressions", "Clicks", "Clicks", "Impressions"]
                values = [100, 10, 20, 200]
                headers.insert(date_position, "Date")
                values.insert(date_position, "2025/1/1")
                self.ws.append(headers)
                self.ws.append(values)
                groups = [[i+1 for i in range(len(headers)) if i != date_position][:2],
                          [i+1 for i in range(len(headers)) if i != date_position][2:]]
                response = {"table_count": 2, "blocks": [
                    {"start_column": min(g), "end_column": max(g), "date_column": date_position + 1,
                     "label_row": 0, "label_column": 0, "confidence": 0.95} for g in groups
                ]}
                with patch("process.media_cleaner.structure.call_ollama_layout", return_value=response):
                    audit, records = self.run_cleaner(llm=True, layout=True)
                self.assertEqual(len(records), 2)
                self.assertEqual([(r["Impressions"], r["Clicks (all)"]) for r in records], [(100, 10), (200, 20)])
                self.assertTrue(all(a.detection_method == "llm-layout" for a in audit))

    def test_llm_receives_full_row_and_can_add_a_table_outside_seed(self):
        from process.media_cleaner.detection import discover_blocks
        self.put(1, 2, ["Date", "Impressions"])
        self.put(2, 2, ["2025/1/1", 100])
        self.put(1, 12, ["Date", "Clicks", "Views"])
        self.put(2, 12, ["2025/1/2", 20, 300])
        # Simulate discovery finding only the first table.
        seeds = discover_blocks(SheetGrid(self.ws))[:1]
        response = {"table_count": 2, "blocks": [
            {"start_column": 2, "end_column": 3, "date_column": 2, "label_row": 0, "label_column": 0, "confidence": 0.95},
            {"start_column": 12, "end_column": 14, "date_column": 12, "label_row": 0, "label_column": 0, "confidence": 0.95},
        ]}
        with patch("process.media_cleaner.engine.discover_blocks", return_value=seeds), patch(
            "process.media_cleaner.structure.call_ollama_layout", return_value=response,
        ) as call:
            _, records = self.run_cleaner(llm=True, layout=True)
        self.assertEqual(len(records), 2)
        band = call.call_args.args[0]
        self.assertEqual(band["last_column"], 14)
        self.assertIn([14, "Views"], next(row["cells"] for row in band["rows"] if row["row"] == 1))
        self.assertEqual(records[-1]["Views"], 300)

    def test_structure_runs_even_if_seed_columns_already_have_known_aliases(self):
        self.ws.append(["Date", "Impressions"])
        self.ws.append(["2025/1/1", 100])
        response = {"table_count": 1, "blocks": [{"start_column": 1, "end_column": 2, "date_column": 1,
                    "label_row": 0, "label_column": 0, "confidence": 0.95}]}
        with patch("process.media_cleaner.structure.call_ollama_layout", return_value=response) as call:
            _, records = self.run_cleaner(llm=True, layout=True)
        self.assertEqual(call.call_count, 1)
        self.assertEqual(len(records), 1)

    def test_invalid_or_incomplete_layouts_fall_back_without_losing_rows(self):
        self.ws.append(["Date", "Impressions", "Clicks"])
        self.ws.append(["2025/1/1", 100, 10])
        valid = {"start_column": 1, "end_column": 3, "date_column": 1, "label_row": 0, "label_column": 0, "confidence": 0.95}
        responses = [
            {"table_count": 2, "blocks": [valid]},
            {"table_count": 1, "blocks": [{**valid, "end_column": 100}]},
            {"table_count": 1, "blocks": [{**valid, "end_column": 2}]},
            {"table_count": 2, "blocks": [valid, valid]},
            {"table_count": 1, "blocks": [{**valid, "date_column": 2}]},
        ]
        for response in responses:
            with self.subTest(response=response), patch("process.media_cleaner.structure.call_ollama_layout", return_value=response):
                audit, records = self.run_cleaner(llm=True, layout=True)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]["Clicks (all)"], 10)
            self.assertTrue(any("LLM 結構提案未採用" in w for w in audit[0].warnings))

    def test_unknown_fields_receive_samples_of_shared_date_outside_region(self):
        self.ws.append(["陌生展示量", "Date"])
        self.ws.append([100, "2025/1/1"])
        layout = {"table_count": 1, "blocks": [{"start_column": 1, "end_column": 1,
                  "date_column": 2, "confidence": 0.95}]}
        mappings = {"blocks": [{"id": "table_1", "mappings": [
            {"column": 1, "target": "Impressions", "confidence": 0.95},
        ]}]}
        with patch("process.media_cleaner.structure.call_ollama_layout", return_value=layout), patch(
            "process.media_cleaner.mapping.call_ollama_json", return_value=mappings,
        ) as model:
            _, records = self.run_cleaner(llm=True, layout=True)
        prompt = model.call_args.args[0]
        self.assertIn('"date_source"', prompt)
        self.assertIn('"2025/1/1"', prompt)
        self.assertEqual(records[0]["Impressions"], 100)

    def test_unmapped_and_empty_source_columns_are_preserved_after_source(self):
        self.ws.append(["Date", "Impressions", "Vendor metric", "Empty metric", "Total"])
        self.ws.append(["2025/1/1", 100, 7, None, 107])
        audit, records = self.run_cleaner()
        self.assertEqual(records[0]["Vendor metric"], 7)
        self.assertIsNone(records[0]["Empty metric"])
        self.assertNotIn("Total", records[0])
        self.assertEqual(records[0]["Source"], "source.xlsx")
        output = load_workbook(self.root / "result.xlsx", data_only=True, read_only=True)
        try:
            headers = list(next(output.active.values))
        finally:
            output.close()
        self.assertEqual(headers, ["Date", "Impressions", "Source", "Vendor metric", "Empty metric"])
        self.assertEqual(audit[0].unmapped_columns, ["Vendor metric", "Empty metric"])

    def test_date_table_with_only_unknown_metric_is_still_exported(self):
        self.ws.append(["Date", "Brand lift score"])
        self.ws.append(["2025/1/1", 42])
        audit, records = self.run_cleaner()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Brand lift score"], 42)
        self.assertEqual(audit[0].status, "成功")

    def test_duplicate_unknown_headers_do_not_overwrite_each_other(self):
        self.ws.append(["Date", "Custom", "Custom"])
        self.ws.append(["2025/1/1", 10, 20])
        audit, records = self.run_cleaner()
        self.assertEqual(records[0]["Custom"], 10)
        self.assertEqual(records[0]["Custom [C]"], 20)
        self.assertTrue(any("重複" in warning for warning in audit[0].warnings))

    def test_source_input_field_and_excel_provenance_are_both_retained(self):
        self.ws.append(["Date", "Source", "Unknown"])
        self.ws.append(["2025/1/1", "Facebook", "x"])
        _, records = self.run_cleaner()
        self.assertEqual(records[0]["Media"], "Facebook")
        self.assertEqual(records[0]["Source"], "source.xlsx")
        self.assertEqual(records[0]["Unknown"], "x")


if __name__ == "__main__":
    unittest.main()
