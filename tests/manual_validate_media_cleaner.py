"""Opt-in validation with a real Ollama model and private local workbooks.

This is deliberately outside unittest discovery. Synthetic variants share values
but change layout, so successful imports must preserve both row counts and sums.
Private workbooks are a development set, not an independent generalization set.
"""

import argparse
import json
from pathlib import Path
from time import monotonic

from openpyxl import Workbook, load_workbook

from process import media_cleaner as cleaner
from process.media_cleaner import structure, mapping
from unittest.mock import patch


def synthetic_cases(folder):
    cases = []
    for position in (0, 2, 4):
        headers = ["Impressions", "Clicks", "Clicks", "Impressions"]
        values = [100, 10, 20, 200]
        headers.insert(position, "Date")
        values.insert(position, "2025/1/1")
        cases.append((f"shared_date_position_{position}", [headers, values], 2, 300, 30))
    cases.append(("different_widths_with_gap", [
        ["Date", "Impressions", None, None, "Views", "Clicks", "Date"],
        ["2025/1/1", 300, None, None, 500, 30, "2025/1/2"],
    ], 2, 300, 30))
    cases.append(("internal_blank_rows", [
        ["Date", "Impressions", "Clicks"], ["2025/1/1", 100, 10], [], [], ["2025/1/2", 200, 20],
    ], 2, 300, 30))
    cases.append(("unknown_header_names", [
        ["出街日", "Ad impressions delivered", "Number of ad clicks"], ["2025/1/1", 300, 30],
    ], 1, 300, 30))
    for name, rows, expected_rows, impressions, clicks in cases:
        path = folder / (name + ".xlsx")
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        workbook.save(path)
        workbook.close()
        yield path, (expected_rows, impressions, clicks)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--source", type=Path)
    parser.add_argument("--model", default="qwen3.5:9b")
    parser.add_argument("--no-llm", action="store_true")
    parser.add_argument("--only", nargs="+", help="只執行檔名包含任一指定文字的案例")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    fixtures = args.output / "synthetic_inputs"
    fixtures.mkdir(exist_ok=True)
    cases = list(synthetic_cases(fixtures))
    if args.source:
        cases.extend((p, None) for p in cleaner.discover_inputs(args.source))
    if args.only:
        cases = [(p, expected) for p, expected in cases if any(name in p.name for name in args.only)]
    aliases = cleaner.read_dictionary(None)
    summaries, audits, paths = [], [], []
    for source, expected in cases:
        started = monotonic()
        output = args.output / (source.stem + "_cleaned.xlsx")
        config = cleaner.OllamaConfig(
            model=args.model, timeout=180, enabled=not args.no_llm,
            force_all=not args.no_llm,
        )
        traces = []
        request_layout = structure.call_ollama_layout
        request_json = mapping.call_ollama_json
        mapping_traces = []
        def trace_json(prompt, config, **kwargs):
            response = request_json(prompt, config, **kwargs)
            if "schema" not in kwargs:
                mapping_traces.append({"prompt": prompt, "response": response})
                (args.output / (source.stem + "_mapping_trace.json")).write_text(json.dumps(mapping_traces, ensure_ascii=False, indent=2))
            return response
        def trace_layout(band, config, feedback=""):
            response = request_layout(band, config, feedback)
            traces.append({"band": band, "feedback": feedback, "response": response})
            (args.output / (source.stem + "_layout_trace.json")).write_text(json.dumps(traces, ensure_ascii=False, indent=2))
            return response
        with patch.object(structure, "call_ollama_layout", side_effect=trace_layout), patch.object(
            mapping, "call_ollama_json", side_effect=trace_json,
        ):
            records, count = cleaner.clean_workbook_sheets(
                source, output, aliases, 0, cleaner.list_workbook_sheets(source.read_bytes()), config,
            )
        metrics = {"Impressions": 0, "Clicks (all)": 0}
        if count:
            book = load_workbook(output, data_only=True, read_only=True)
            values = iter(book.active.values)
            headers = next(values)
            for values_row in values:
                record = dict(zip(headers, values_row))
                for key in metrics:
                    value = record.get(key)
                    if isinstance(value, (int, float)):
                        metrics[key] += value
            book.close()
            paths.append(output)
        observed = (count, metrics["Impressions"], metrics["Clicks (all)"])
        summary = {
            "file": source.name, "rows": count, "metric_sums": metrics,
            "expected": expected, "passed": observed == expected if expected else None,
            "llm_layout_blocks": sum(r.detection_method.startswith("llm-layout") for r in records),
            "warnings": sorted({warning for r in records for warning in r.warnings}),
            "seconds": round(monotonic() - started, 1),
        }
        summaries.append(summary)
        audits.extend(records)
        print(json.dumps(summary, ensure_ascii=False), flush=True)
        (args.output / "validation_summary.json").write_text(json.dumps(summaries, ensure_ascii=False, indent=2))
        cleaner.write_audit(audits, args.output)
    cleaner.consolidate_cleaned_workbooks(paths, audits, args.output / "cleaned_media_results.xlsx")
    return 1 if any(s["passed"] is False for s in summaries) else 0


if __name__ == "__main__":
    raise SystemExit(main())
